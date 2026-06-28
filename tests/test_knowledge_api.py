import io
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pypdf import PdfWriter
from starlette.datastructures import UploadFile

from app.main import app


client = TestClient(app)


def _configure_knowledge(monkeypatch, tmp_path: Path) -> None:
    monkeypatch.setenv("KNOWLEDGE_ENABLED", "true")
    monkeypatch.setenv("KNOWLEDGE_STORE_FILE", str(tmp_path / "knowledge" / "index.json"))
    monkeypatch.setenv("KNOWLEDGE_UPLOAD_DIR", str(tmp_path / "knowledge" / "uploads"))
    monkeypatch.setenv("KNOWLEDGE_ALLOWED_DIRS", str(tmp_path / "allowed"))
    monkeypatch.setenv("KNOWLEDGE_MAX_UPLOAD_MB", "1")
    monkeypatch.setenv("KNOWLEDGE_CHUNK_SIZE", "40")
    monkeypatch.setenv("KNOWLEDGE_CHUNK_OVERLAP", "5")


def _upload(files: list[tuple[str, tuple[str, bytes, str]]]):
    return client.post("/assistant/knowledge/upload", files=files)


def _minimal_pdf() -> bytes:
    output = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.write(output)
    return output.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("Lokaler Vertrag")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _xlsx_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["Kategorie", "Betrag"])
    worksheet.append(["Boden", 28.9])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def test_upload_indexes_supported_text_file_and_returns_safe_document(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)

    response = _upload([("files", ("notizen.txt", b"Lokales Wissen zum Hauskauf", "text/plain"))])

    assert response.status_code == 200
    payload = response.json()
    assert payload["uploaded"] is True
    assert payload["success_count"] == 1
    document = payload["documents"][0]
    assert document["name"] == "notizen.txt"
    assert document["chunk_count"] >= 1
    assert document["path"]
    assert "text" not in document


def test_upload_processes_each_file_independently_and_keeps_partial_success(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)

    response = _upload(
        [
            ("files", ("gueltig.txt", b"Inhalt", "text/plain")),
            ("files", ("unbekannt.exe", b"nicht erlaubt", "application/octet-stream")),
        ]
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["uploaded"] is True
    assert payload["success_count"] == 1
    assert payload["failed_count"] == 1
    assert payload["errors"][0]["reason"] == "unsupported_file_type"


def test_upload_deduplicates_without_creating_or_reindexing_a_second_document(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)

    first = _upload([("files", ("eins.txt", b"Gleicher Inhalt", "text/plain"))])
    second = _upload([("files", ("zwei.txt", b"Gleicher Inhalt", "text/plain"))])

    assert first.status_code == 200
    assert second.status_code == 200
    assert second.json()["documents"][0]["duplicate"] is True
    assert client.get("/assistant/knowledge/documents").json()["count"] == 1
    upload_dir = Path(tmp_path / "knowledge" / "uploads")
    assert len([path for path in upload_dir.iterdir() if path.is_file()]) == 1


def test_upload_rejects_empty_large_traversal_and_bad_pdf_safely(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    too_large = b"x" * (1024 * 1024 + 1)

    response = _upload(
        [
            ("files", ("leer.txt", b"", "text/plain")),
            ("files", ("../traversal.txt", b"no", "text/plain")),
            ("files", ("gross.txt", too_large, "text/plain")),
            ("files", ("schlecht.pdf", b"not a pdf", "application/pdf")),
        ]
    )

    assert response.status_code == 200
    errors = {item["filename"]: item for item in response.json()["errors"]}
    assert errors["leer.txt"]["reason"] == "empty_file"
    assert errors["traversal.txt"]["reason"] == "invalid_filename"
    assert errors["gross.txt"]["reason"] == "file_too_large"
    assert errors["schlecht.pdf"]["reason"] == "invalid_pdf_header"
    assert all("exception" not in str(item).lower() for item in errors.values())


def test_upload_supports_pdf_docx_xlsx_csv_and_json(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)

    response = _upload(
        [
            ("files", ("leer.pdf", _minimal_pdf(), "application/pdf")),
            ("files", ("vertrag.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")),
            ("files", ("kosten.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
            ("files", ("werte.csv", b"Kategorie;Betrag\nBoden;28,90\n", "text/csv")),
            ("files", ("daten.json", b'{"lokal": true}', "application/json")),
        ]
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["success_count"] == 5
    by_name = {item["name"]: item for item in payload["documents"]}
    assert by_name["leer.pdf"]["extraction_status"] == "ocr_required"
    assert by_name["vertrag.docx"]["chunk_count"] >= 1
    assert by_name["kosten.xlsx"]["chunk_count"] >= 1
    pdf_id = by_name["leer.pdf"]["document_id"]
    assert client.get("/assistant/knowledge/documents").json()["documents"][0]["extraction_status"] == "ocr_required"
    detail = client.get(f"/assistant/knowledge/documents/{pdf_id}").json()
    assert detail["document"]["extraction_status"] == "ocr_required"
    assert detail["document"]["extraction_message"] == "Das PDF enthält keinen extrahierbaren Text. OCR wird noch nicht unterstützt."
    reindexed = client.post(f"/assistant/knowledge/documents/{pdf_id}/reindex")
    assert reindexed.status_code == 200
    assert reindexed.json()["indexed"] is True
    assert reindexed.json()["document"]["extraction_status"] == "ocr_required"


def test_document_detail_is_bounded_and_does_not_expose_full_content(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    body = ("vertraulicher lokaler Text " * 200).encode("utf-8")
    uploaded = _upload([("files", ("lang.txt", body, "text/plain"))]).json()
    document_id = uploaded["documents"][0]["document_id"]

    response = client.get(f"/assistant/knowledge/documents/{document_id}")

    assert response.status_code == 200
    detail = response.json()
    assert detail["document"]["document_id"] == document_id
    assert detail["document"]["path"]
    assert len(detail["chunks"]) <= 5
    assert all(len(chunk["preview"]) <= 320 for chunk in detail["chunks"])
    assert all(chunk["chunk_id"] for chunk in detail["chunks"])
    assert all(isinstance(chunk["index"], int) for chunk in detail["chunks"])
    assert "vertraulicher lokaler Text " * 50 not in str(detail)


def test_document_detail_reindex_and_delete_have_safe_not_found_behavior(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    missing = "does-not-exist"

    assert client.get(f"/assistant/knowledge/documents/{missing}").status_code == 404
    assert client.post(f"/assistant/knowledge/documents/{missing}/reindex").status_code == 404
    assert client.delete(f"/assistant/knowledge/documents/{missing}").status_code == 404


def test_reindex_and_delete_uploaded_document_returns_metadata_path(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    uploaded = _upload([("files", ("delete.txt", b"Bitte erneut indexieren", "text/plain"))]).json()
    document_id = uploaded["documents"][0]["document_id"]

    reindexed = client.post(f"/assistant/knowledge/documents/{document_id}/reindex")
    deleted = client.delete(f"/assistant/knowledge/documents/{document_id}")

    assert reindexed.status_code == 200
    assert reindexed.json()["document"]["document_id"] == document_id
    assert reindexed.json()["document"]["path"]
    assert deleted.status_code == 200
    assert deleted.json()["deleted"] is True
    assert deleted.json()["document_id"] == document_id


def test_reindex_replaces_existing_document_chunks(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    uploaded = _upload([("files", ("ersetzen.txt", b"Alter Inhalt", "text/plain"))]).json()
    document_id = uploaded["documents"][0]["document_id"]
    stored_path = Path(client.get("/assistant/knowledge/documents").json()["documents"][0]["path"])
    stored_path.write_text("Neuer Inhalt mit Energieausweis", encoding="utf-8")

    response = client.post(f"/assistant/knowledge/documents/{document_id}/reindex")
    detail = client.get(f"/assistant/knowledge/documents/{document_id}").json()

    assert response.status_code == 200
    assert response.json()["indexed"] is True
    assert "Neuer Inhalt" in detail["chunks"][0]["preview"]
    assert "Alter Inhalt" not in detail["chunks"][0]["preview"]


def test_delete_local_path_document_keeps_original_file(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    source = tmp_path / "allowed" / "vertrag.txt"
    source.parent.mkdir(parents=True)
    source.write_text("Kaufvertrag", encoding="utf-8")
    indexed = client.post("/assistant/knowledge/index", json={"path": str(source)}).json()

    response = client.delete(f"/assistant/knowledge/documents/{indexed['document']['document_id']}")

    assert response.status_code == 200
    assert response.json()["physical_file_deleted"] is False
    assert source.exists()


def test_reindex_reports_missing_uploaded_source_with_safe_conflict(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    uploaded = _upload([("files", ("missing.txt", b"Lokaler Inhalt", "text/plain"))]).json()
    document_id = uploaded["documents"][0]["document_id"]
    stored_path = Path(client.get("/assistant/knowledge/documents").json()["documents"][0]["path"])
    stored_path.unlink()

    response = client.post(f"/assistant/knowledge/documents/{document_id}/reindex")

    assert response.status_code == 409
    assert response.json()["detail"]["reason"] == "source_file_missing"
    assert "exception" not in response.text.lower()


def test_partial_multi_upload_closes_each_upload_file(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    closed: list[str | None] = []
    original_close = UploadFile.close

    async def close_spy(self):
        closed.append(self.filename)
        await original_close(self)

    monkeypatch.setattr(UploadFile, "close", close_spy)

    response = _upload(
        [
            ("files", ("geschlossen.txt", b"Inhalt", "text/plain")),
            ("files", ("abgelehnt.exe", b"Inhalt", "application/octet-stream")),
        ]
    )

    assert response.status_code == 200
    assert "geschlossen.txt" in closed
    assert "abgelehnt.exe" in closed


def test_status_documents_search_compatibility_and_no_secrets(monkeypatch, tmp_path: Path) -> None:
    _configure_knowledge(monkeypatch, tmp_path)
    monkeypatch.setenv("SHOULD_NOT_APPEAR_SECRET", "not-a-secret-value")
    _upload([("files", ("suche.md", b"Hammer Jarvis bleibt lokal", "text/markdown"))])

    status = client.get("/assistant/knowledge/status")
    documents = client.get("/assistant/knowledge/documents")
    search = client.get("/assistant/knowledge/search?q=lokal")

    assert status.status_code == 200
    assert status.json()["embedding_enabled"] is False
    assert status.json()["search_mode"] == "keyword"
    assert status.json()["max_upload_mb"] == 1
    assert status.json()["extractor_status"] == {
        "pdf": True,
        "docx": True,
        "xlsx": True,
        "xlsm": True,
        "csv": True,
        "txt": True,
        "md": True,
        "json": True,
        "ocr": False,
    }
    assert ".txt" in status.json()["supported_extensions"]
    assert status.json()["data_dir"] == str(tmp_path / "knowledge")
    assert status.json()["upload_dir"] == str(tmp_path / "knowledge" / "uploads")
    assert "not-a-secret-value" not in status.text
    assert documents.status_code == 200
    assert all("text" not in item for item in documents.json()["documents"])
    assert all("chunks" not in item for item in documents.json()["documents"])
    assert search.status_code == 200
    assert search.json()["count"] == 1
    assert {"data_dir", "upload_dir", "store_file", "document_count", "chunk_count", "total_size_bytes", "supported_extensions", "max_upload_mb", "embedding_enabled", "search_mode", "extractor_status"}.issubset(status.json())
    document = documents.json()["documents"][0]
    assert {"document_id", "name", "path", "extension", "mime_type", "size_bytes", "chunk_count", "extraction_status", "source_type"}.issubset(document)
    assert "chunks" not in document
    assert "text" not in document
